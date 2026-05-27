import io
from datetime import datetime
from flask import Blueprint, jsonify, request, g, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from models import db, Customer, GeneratedEmail
from middleware import login_required

export_bp = Blueprint("export", __name__)

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Arial", size=10)
CUSTOMER_HEADERS = [
    "公司名称", "国家", "联系人", "邮箱", "电话", "网址",
    "主营产品", "客户来源", "跟进状态", "下次跟进时间", "备注",
]


def _build_customer_query(uid, mode, ids, filters):
    """Build a customer query based on export mode."""
    stmt = db.select(Customer).where(Customer.user_id == uid)

    if mode == "selected" and ids:
        stmt = stmt.where(Customer.id.in_(ids))
    elif mode == "filtered" and filters:
        q = (filters.get("q") or "").strip()
        status = (filters.get("status") or "").strip()
        favorite = (filters.get("favorite") or "").strip()
        if q:
            like = f"%{q}%"
            stmt = stmt.where(
                db.or_(
                    Customer.company_name.ilike(like),
                    Customer.email.ilike(like),
                    Customer.country.ilike(like),
                    Customer.contact_name.ilike(like),
                )
            )
        if status:
            stmt = stmt.where(Customer.status == status)
        if favorite == "true":
            stmt = stmt.where(Customer.favorite == True)

    return stmt.order_by(Customer.updated_at.desc())


def _format_date(dt):
    """Format a datetime or date to YYYY-MM-DD string."""
    if not dt:
        return ""
    if hasattr(dt, "strftime"):
        return dt.strftime("%Y-%m-%d")
    return str(dt)[:10]


@export_bp.post("/customers/xlsx")
@login_required
def export_customers_xlsx():
    """Export customers to Excel (.xlsx)."""
    uid = g.current_user.id
    payload = request.get_json(silent=True) or {}
    mode = (payload.get("mode") or "all").strip()
    ids = [int(i) for i in (payload.get("ids") or []) if str(i).isdigit()]
    filters = payload.get("filters") or {}

    if mode not in ("all", "filtered", "selected"):
        return jsonify({"error": "Invalid mode. Allowed: all, filtered, selected."}), 400

    stmt = _build_customer_query(uid, mode, ids, filters)
    customers = db.session.execute(stmt).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "客户列表"

    # Header row
    for col, header in enumerate(CUSTOMER_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    status_labels = {"potential": "潜在客户", "contacted": "已联系", "following": "跟进中", "won": "已成交", "lost": "无效客户"}
    for row, c in enumerate(customers, 2):
        row_data = [
            c.company_name, c.country, c.contact_name, c.email, c.phone,
            c.website, c.product_keyword, c.source,
            status_labels.get(c.status, c.status),
            _format_date(c.next_follow_up), c.notes,
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = CELL_FONT
            cell.alignment = Alignment(vertical="center")

    # Auto-width
    for col in range(1, len(CUSTOMER_HEADERS) + 1):
        max_width = max(len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, len(customers) + 2))
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(max_width + 4, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    today = datetime.now().strftime("%Y%m%d")
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"customers_export_{today}.xlsx",
    )


@export_bp.post("/emails/word")
@login_required
def export_emails_word():
    """Export saved emails to Word (.docx)."""
    uid = g.current_user.id
    payload = request.get_json(silent=True) or {}
    ids = [int(i) for i in (payload.get("ids") or []) if str(i).isdigit()]

    if not ids:
        return jsonify({"error": "No email IDs provided."}), 400

    emails = db.session.execute(
        db.select(GeneratedEmail).where(
            GeneratedEmail.id.in_(ids),
            GeneratedEmail.user_id == uid,
        ).order_by(GeneratedEmail.created_at.desc())
    ).scalars().all()

    if not emails:
        return jsonify({"error": "No emails found."}), 404

    doc = Document()

    # Title
    title = doc.add_heading("AI Generated Emails", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    scene_labels = {
        "first_contact": "首次开发信", "sample_followup": "样品跟进",
        "order_confirmation": "订单确认", "holiday_greeting": "节日问候",
        "after_sales": "售后跟进",
    }

    for i, email in enumerate(emails):
        if i > 0:
            doc.add_page_break()

        scene_label = scene_labels.get(email.scene, email.scene)
        heading_text = f"{email.company_name} — {email.product_name}"
        if not heading_text.strip().replace("—", "").strip():
            heading_text = f"Email #{email.id}"

        doc.add_heading(heading_text, level=1)

        # Meta info
        meta = doc.add_paragraph()
        meta.style = doc.styles["Normal"]
        meta_parts = []
        if email.country:
            meta_parts.append(f"Country: {email.country}")
        if email.industry:
            meta_parts.append(f"Industry: {email.industry}")
        meta_parts.append(f"Scene: {scene_label}")
        meta_parts.append(f"Date: {_format_date(email.created_at)}")
        meta.add_run(" | ".join(meta_parts)).font.size = Pt(9)

        doc.add_paragraph()  # spacer

        # Email body
        for line in email.body.split("\n"):
            p = doc.add_paragraph(line)
            p.style = doc.styles["Normal"]
            for run in p.runs:
                run.font.size = Pt(11)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)

    today = datetime.now().strftime("%Y%m%d")
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=f"emails_export_{today}.docx",
    )
